"""Files that need no OCR: Office documents and plain-text formats, converted straight to Markdown."""

from __future__ import annotations

import csv
import io
import itertools
from pathlib import Path

TEXT_EXTS = {".txt", ".md", ".markdown"}
TABLE_EXTS = {".csv", ".tsv"}
HTML_EXTS = {".html", ".htm", ".xhtml"}
# Data and code files are kept verbatim in a fenced block, tagged with their language.
CODE_EXTS = {
    ".json": "json", ".xml": "xml", ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".ini": "ini",
    ".log": "text", ".sql": "sql", ".py": "python", ".js": "javascript", ".ts": "typescript", ".sh": "bash",
}
OFFICE_XML_EXTS = {".docx", ".pptx", ".xlsx"}
DIRECT_EXTS = TEXT_EXTS | TABLE_EXTS | HTML_EXTS | set(CODE_EXTS) | OFFICE_XML_EXTS


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def markdown_table(rows: list[list]) -> str:
    def cell(value) -> str:
        return "" if value is None else str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>").strip()

    rows = [[cell(v) for v in row] for row in rows]
    width = max((len(r) for r in rows), default=0)
    if not width:
        return ""
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _trim(rows: list[list]) -> list[list]:
    """Drop empty rows and the empty columns on the right (spreadsheets report a used range)."""
    rows = [list(r) for r in rows if any(v not in (None, "") for v in r)]
    width = max((max((i + 1 for i, v in enumerate(r) if v not in (None, "")), default=0) for r in rows), default=0)
    return [r[:width] for r in rows]


def _image_saver(images_dir: Path, stem: str):
    counter = itertools.count(1)

    def save(data: bytes, ext: str) -> str:
        images_dir.mkdir(parents=True, exist_ok=True)
        name = f"{stem}_img{next(counter)}.{ext.lower().replace('jpeg', 'jpg')}"
        (images_dir / name).write_bytes(data)
        return f"{images_dir.name}/{name}"

    return save


def docx_to_markdown(path: Path, images_dir: Path, stem: str) -> str:
    import mammoth

    save = _image_saver(images_dir, stem)

    def convert_image(image):
        with image.open() as data:
            return {"src": save(data.read(), image.content_type.split("/")[-1])}

    style_map = "p[style-name='Title'] => h1:fresh\np[style-name='Subtitle'] => h2:fresh"
    with open(path, "rb") as f:
        html = mammoth.convert_to_html(
            f, convert_image=mammoth.images.img_element(convert_image), style_map=style_map
        ).value
    return html_to_markdown(html)


def html_to_markdown(html: str) -> str:
    from markdownify import markdownify

    return markdownify(html, heading_style="ATX", bullets="-", table_infer_header=True)


def pptx_to_markdown(path: Path, images_dir: Path, stem: str) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    save = _image_saver(images_dir, stem)

    def shape_markdown(shape) -> list[str]:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return [md for child in shape.shapes for md in shape_markdown(child)]
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return [f"![]({save(shape.image.blob, shape.image.ext)})"]
        if getattr(shape, "has_table", False) and shape.has_table:
            return [markdown_table([[c.text for c in row.cells] for row in shape.table.rows])]
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            paragraphs = [(p.level, "".join(r.text for r in p.runs).strip()) for p in shape.text_frame.paragraphs]
            paragraphs = [(level, text) for level, text in paragraphs if text]
            if len(paragraphs) == 1 and paragraphs[0][0] == 0:
                return [paragraphs[0][1]]
            return ["\n".join("  " * level + "- " + text for level, text in paragraphs)] if paragraphs else []
        return []

    parts = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        parts.append(f"<!-- slide {number} -->\n\n## {title or f'Slide {number}'}")
        shapes = sorted((s for s in slide.shapes if s != title_shape), key=lambda s: (s.top or 0, s.left or 0))
        parts += [md for s in shapes for md in shape_markdown(s)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes = slide.notes_slide.notes_text_frame.text.strip().replace("\n", "\n> ")
            parts.append(f"> **Notes:** {notes}")
    return "\n\n".join(parts)


def xlsx_to_markdown(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows = _trim(list(sheet.iter_rows(values_only=True)))
        parts.append(f"## {sheet.title}\n\n" + (markdown_table(rows) if rows else "_(empty sheet)_"))
    workbook.close()
    return "\n\n".join(parts)


def table_file_to_markdown(path: Path) -> str:
    text = read_text(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return markdown_table(_trim(list(csv.reader(io.StringIO(text), delimiter=delimiter))))


def direct_to_markdown(path: Path, images_dir: Path) -> str:
    """Markdown for a file in DIRECT_EXTS."""
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return read_text(path)
    if ext in TABLE_EXTS:
        return table_file_to_markdown(path)
    if ext in HTML_EXTS:
        return html_to_markdown(read_text(path))
    if ext in CODE_EXTS:
        text = read_text(path).rstrip("\n")
        fence = "````" if "```" in text else "```"
        return f"{fence}{CODE_EXTS[ext]}\n{text}\n{fence}"
    if ext == ".docx":
        return docx_to_markdown(path, images_dir, path.stem)
    if ext == ".pptx":
        return pptx_to_markdown(path, images_dir, path.stem)
    if ext == ".xlsx":
        return xlsx_to_markdown(path)
    raise ValueError(f"{path.name}: not a directly convertible file")
