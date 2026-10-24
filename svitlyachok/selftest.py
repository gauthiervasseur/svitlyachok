"""Headless checks of a built app, for CI and support: `Svitlyachok --selftest [--ocr]`.

Writes selftest.log to the current directory (a windowed .exe has no console) and returns
exit code 0 when every check passes. `--ocr` also downloads the model and OCRs a rendered page.
"""

from __future__ import annotations

import sys
import tempfile
import time
import traceback
from pathlib import Path

SAMPLE_HTML = (
    "<h1>Firefly Report</h1><p>The quick brown fox jumps over the lazy dog.</p>"
    "<table border='1'><tr><td>North</td><td>1480</td></tr></table>"
)


def _sample_pdf(path: Path) -> Path:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_htmlbox(pymupdf.Rect(60, 60, 540, 400), SAMPLE_HTML)
    doc.save(path)
    return path


def _check_engine_imports(backend_name: str) -> str:
    if backend_name == "mlx":
        import mlx.core as mx
        import mlx_vlm.models.unlimited_ocr  # noqa: F401  (loaded dynamically by mlx_vlm.load)

        assert mx.array([1, 2]).sum(stream=mx.cpu).item() == 3  # CI runners may have no Metal GPU
        return f"mlx {'with' if mx.metal.is_available() else 'WITHOUT'} Metal GPU"
    import torch

    # Imported at runtime by the model's own code (trust_remote_code), so the bundle must ship them.
    import addict, easydict, einops, matplotlib.pyplot, requests, torchvision.transforms  # noqa: E401,F401
    import transformers.models.llama.modeling_llama  # noqa: F401

    assert torch.ones(2).sum().item() == 2
    return f"torch {torch.__version__}"


def run(ocr: bool = False) -> int:
    from .app import UI_DIR
    from .engine.core import UnlimitedOCR

    lines: list[str] = []
    passed = True

    def check(name, fn):
        nonlocal passed
        start = time.perf_counter()
        try:
            detail = fn() or ""
            lines.append(f"PASS {name} ({time.perf_counter() - start:.1f}s) {detail}")
        except Exception:
            passed = False
            lines.append(f"FAIL {name}\n{traceback.format_exc()}")

    engine = UnlimitedOCR()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        def ui():
            assert (UI_DIR / "index.html").exists() and (UI_DIR / "fonts" / "PressStart2P-Regular.ttf").exists()

        def webview_import():
            import webview  # noqa: F401
            from webview.dom import DOMEventHandler  # noqa: F401

        def text_layer():
            result = engine.convert(_sample_pdf(tmp / "sample.pdf"), tmp / "pdf")
            text = result.md_path.read_text(encoding="utf-8")
            assert "Firefly Report" in text and result.ocr_pages == 0, text
            return f"{result.text_pages} page(s) read without OCR"

        def office():
            import openpyxl
            from pptx import Presentation

            book = openpyxl.Workbook()
            book.active.append(["Region", "Q3"])
            book.active.append(["North", 1480])
            book.save(tmp / "sheet.xlsx")
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Firefly deck"
            deck.save(tmp / "deck.pptx")
            assert "| North | 1480 |" in engine.convert(tmp / "sheet.xlsx", tmp / "x").md_path.read_text(encoding="utf-8")
            assert "Firefly deck" in engine.convert(tmp / "deck.pptx", tmp / "p").md_path.read_text(encoding="utf-8")

        def ocr_page():
            import pymupdf

            png = tmp / "scan.png"
            with pymupdf.open(_sample_pdf(tmp / "scan.pdf")) as doc:
                doc[0].get_pixmap(dpi=200).save(png)
            result = engine.convert(png, tmp / "ocr")
            text = result.md_path.read_text(encoding="utf-8")
            assert "firefly" in text.lower() and result.ocr_pages == 1, text
            return f"{engine.backend.name}, model load {engine.load_seconds:.0f}s"

        check("ui files", ui)
        check("webview", webview_import)
        check("engine libraries", lambda: _check_engine_imports(engine.backend_name))
        check("pdf text layer", text_layer)
        check("office files", office)
        if ocr:
            check("ocr", ocr_page)

    report = "\n".join(lines) + ("\nALL PASSED\n" if passed else "\nFAILED\n")
    Path("selftest.log").write_text(report, encoding="utf-8")
    print(report)
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(run(ocr="--ocr" in sys.argv))
